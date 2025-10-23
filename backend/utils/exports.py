from io import BytesIO
from datetime import datetime, timedelta

from quart import Response
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


async def generate_pdf_statement(account_summary: dict, trades: list):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    
    # Container for the 'Flowable' objects
    elements = []
    
    # Define styles
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a237e'),
        spaceAfter=30,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#283593'),
        spaceAfter=12,
        spaceBefore=12,
        fontName='Helvetica-Bold'
    )
    
    normal_style = ParagraphStyle(
        'CustomNormal',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#424242')
    )
    
    # Title
    title = Paragraph("<b>TRADING STATEMENT</b>", title_style)
    elements.append(title)
    
    # Generated date and license info
    date_text = f"Generated: {datetime.utcnow().strftime('%B %d, %Y at %H:%M:%S')} UTC"
    elements.append(Paragraph(date_text, normal_style))
    
    if 'license_key' in account_summary:
        license_text = f"License Key: <b>{account_summary['license_key']}</b>"
        elements.append(Paragraph(license_text, normal_style))
    
    elements.append(Spacer(1, 0.3*inch))
    
    # Account Summary Section
    elements.append(Paragraph("Account Summary", heading_style))
    
    account_data = [
        ['Metric', 'Value'],
        ['Balance', f"${account_summary.get('balance', 0):.2f}"],
        ['Equity', f"${account_summary.get('equity', 0):.2f}"],
    ]
    
    account_table = Table(account_data, colWidths=[3*inch, 3*inch])
    account_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1565c0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TOPPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#e3f2fd')),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#90caf9')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#e3f2fd'), colors.white]),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    elements.append(account_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Performance Metrics Section
    elements.append(Paragraph("Performance Metrics (Last 30 Days)", heading_style))
    
    win_rate = account_summary.get('win_rate', 0)
    win_rate_color = colors.HexColor('#4caf50') if win_rate >= 50 else colors.HexColor('#f44336')
    
    perf_data = [
        ['Metric', 'Value'],
        ['Win Rate', f"{win_rate:.2f}%"],
        ['Max Drawdown', f"{account_summary.get('drawdown', 0):.2f}%"],
        ['Total Trades', str(account_summary.get('total_trades', 0))],
    ]
    
    perf_table = Table(perf_data, colWidths=[3*inch, 3*inch])
    perf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1565c0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TOPPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#e8f5e9')),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#a5d6a7')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#e8f5e9'), colors.white]),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    elements.append(perf_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Trades Section
    elements.append(Paragraph("Trade History (Last 30 Days)", heading_style))
    
    # Prepare trade data
    trade_data = [['Pair', 'Lots', 'Direction', 'Result', 'Status', 'Closed At']]
    
    for t in trades:
        result = t.get('result', 0)
        result_str = f"${result:.2f}"
        
        trade_data.append([
            t.get('pair', 'N/A'),
            f"{t.get('lots', 0):.2f}",
            t.get('direction', 'N/A'),
            result_str,
            t.get('status', 'N/A').upper(),
            t.get('closed_at', 'N/A')[:10] if t.get('closed_at') else 'Open'
        ])
    
    # Create table with appropriate column widths
    col_widths = [1.2*inch, 0.8*inch, 1*inch, 1*inch, 0.8*inch, 1.2*inch]
    trade_table = Table(trade_data, colWidths=col_widths, repeatRows=1)
    
    # Build style commands
    style_commands = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1565c0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('TOPPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#90caf9')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f5f5f5'), colors.white]),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
    ]
    
    # Color code results (profit = green, loss = red)
    for i, t in enumerate(trades, start=1):
        result = t.get('result', 0)
        if result > 0:
            style_commands.append(('TEXTCOLOR', (3, i), (3, i), colors.HexColor('#2e7d32')))
            style_commands.append(('FONTNAME', (3, i), (3, i), 'Helvetica-Bold'))
        elif result < 0:
            style_commands.append(('TEXTCOLOR', (3, i), (3, i), colors.HexColor('#c62828')))
            style_commands.append(('FONTNAME', (3, i), (3, i), 'Helvetica-Bold'))
    
    trade_table.setStyle(TableStyle(style_commands))
    elements.append(trade_table)
    
    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    
    # Generate filename with license key
    license_key_prefix = account_summary.get('license_key', 'unknown')[:8]
    filename = f'statement_{license_key_prefix}_{datetime.utcnow().date()}.pdf'
    
    # Read the buffer content
    pdf_data = buffer.getvalue()
    
    # Return response with proper headers
    response = Response(pdf_data, mimetype='application/pdf')
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.headers['Content-Length'] = str(len(pdf_data))
    return response


async def generate_xls_statement(account_summary: dict, trades: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "Trading Statement"
    
    # Define colors
    header_color = "1565C0"  # Blue
    alt_row_color = "E3F2FD"  # Light Blue
    perf_header_color = "2E7D32"  # Green
    perf_alt_row_color = "E8F5E9"  # Light Green
    
    # Define border style
    thin_border = Border(
        left=Side(style='thin', color='90CAF9'),
        right=Side(style='thin', color='90CAF9'),
        top=Side(style='thin', color='90CAF9'),
        bottom=Side(style='thin', color='90CAF9')
    )
    
    # Title
    ws.merge_cells('A1:F1')
    title_cell = ws['A1']
    title_cell.value = "TRADING STATEMENT"
    title_cell.font = Font(name='Calibri', size=22, bold=True, color="1A237E")
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 35
    
    # Generated date
    ws.merge_cells('A2:F2')
    date_cell = ws['A2']
    date_cell.value = f"Generated: {datetime.utcnow().strftime('%B %d, %Y at %H:%M:%S')} UTC"
    date_cell.font = Font(name='Calibri', size=10, color="424242")
    date_cell.alignment = Alignment(horizontal='center')
    
    # License Key
    row = 3
    if 'license_key' in account_summary:
        ws.merge_cells(f'A{row}:F{row}')
        license_cell = ws[f'A{row}']
        license_cell.value = f"License Key: {account_summary['license_key']}"
        license_cell.font = Font(name='Calibri', size=10, bold=True, color="424242")
        license_cell.alignment = Alignment(horizontal='center')
        row += 1
    
    row += 1  # Empty row
    
    # Account Summary Section
    ws.merge_cells(f'A{row}:B{row}')
    section_cell = ws[f'A{row}']
    section_cell.value = "Account Summary"
    section_cell.font = Font(name='Calibri', size=14, bold=True, color="283593")
    section_cell.alignment = Alignment(horizontal='left', vertical='center')
    row += 1
    
    # Account Summary Header
    ws[f'A{row}'] = "Metric"
    ws[f'B{row}'] = "Value"
    for col in ['A', 'B']:
        cell = ws[f'{col}{row}']
        cell.fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        cell.font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    row += 1
    
    # Account Summary Data
    account_data = [
        ['Balance', f"${account_summary.get('balance', 0):.2f}"],
        ['Equity', f"${account_summary.get('equity', 0):.2f}"],
    ]
    
    for i, (metric, value) in enumerate(account_data):
        ws[f'A{row}'] = metric
        ws[f'B{row}'] = value
        
        # Apply styling
        for col in ['A', 'B']:
            cell = ws[f'{col}{row}']
            cell.font = Font(name='Calibri', size=10)
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.border = thin_border
            if i % 2 == 0:
                cell.fill = PatternFill(start_color=alt_row_color, end_color=alt_row_color, fill_type="solid")
        
        # Bold the metric name
        ws[f'A{row}'].font = Font(name='Calibri', size=10, bold=True)
        row += 1
    
    row += 1  # Empty row
    
    # Performance Metrics Section
    ws.merge_cells(f'A{row}:B{row}')
    section_cell = ws[f'A{row}']
    section_cell.value = "Performance Metrics (Last 30 Days)"
    section_cell.font = Font(name='Calibri', size=14, bold=True, color="283593")
    section_cell.alignment = Alignment(horizontal='left', vertical='center')
    row += 1
    
    # Performance Header
    ws[f'A{row}'] = "Metric"
    ws[f'B{row}'] = "Value"
    for col in ['A', 'B']:
        cell = ws[f'{col}{row}']
        cell.fill = PatternFill(start_color=perf_header_color, end_color=perf_header_color, fill_type="solid")
        cell.font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    row += 1
    
    # Performance Data
    win_rate = account_summary.get('win_rate', 0)
    perf_data = [
        ['Win Rate', f"{win_rate:.2f}%"],
        ['Max Drawdown', f"{account_summary.get('drawdown', 0):.2f}%"],
        ['Total Trades', str(account_summary.get('total_trades', 0))],
    ]
    
    for i, (metric, value) in enumerate(perf_data):
        ws[f'A{row}'] = metric
        ws[f'B{row}'] = value
        
        # Apply styling
        for col in ['A', 'B']:
            cell = ws[f'{col}{row}']
            cell.font = Font(name='Calibri', size=10)
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.border = thin_border
            if i % 2 == 0:
                cell.fill = PatternFill(start_color=perf_alt_row_color, end_color=perf_alt_row_color, fill_type="solid")
        
        # Bold the metric name
        ws[f'A{row}'].font = Font(name='Calibri', size=10, bold=True)
        
        # Color code win rate
        if metric == 'Win Rate':
            color = "2E7D32" if win_rate >= 50 else "C62828"
            ws[f'B{row}'].font = Font(name='Calibri', size=10, bold=True, color=color)
        
        row += 1
    
    row += 2  # Empty rows
    
    # Trade History Section
    ws.merge_cells(f'A{row}:H{row}')
    section_cell = ws[f'A{row}']
    section_cell.value = "Trade History (Last 30 Days)"
    section_cell.font = Font(name='Calibri', size=14, bold=True, color="283593")
    section_cell.alignment = Alignment(horizontal='left', vertical='center')
    row += 1
    
    # Trade History Header
    headers = ["Pair", "Lots", "Direction", "Result", "Status", "Opened At", "Closed At", "AI Confidence"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx)
        cell.value = header
        cell.fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        cell.font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    
    ws.row_dimensions[row].height = 25
    row += 1
    
    # Trade History Data
    for i, t in enumerate(trades):
        result = t.get('result', 0)
        
        ws.cell(row=row, column=1).value = t.get('pair', 'N/A')
        ws.cell(row=row, column=2).value = t.get('lots', 0)
        ws.cell(row=row, column=3).value = t.get('direction', 'N/A')
        ws.cell(row=row, column=4).value = result
        ws.cell(row=row, column=5).value = t.get('status', 'N/A').upper()
        ws.cell(row=row, column=6).value = t.get('opened_at', 'N/A')[:19] if t.get('opened_at') else 'N/A'
        ws.cell(row=row, column=7).value = t.get('closed_at', 'Open')[:19] if t.get('closed_at') else 'Open'
        ws.cell(row=row, column=8).value = t.get('ai_confidence', 'Unavailable')
        
        # Apply styling to all cells in the row
        for col_idx in range(1, 9):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = Font(name='Calibri', size=9)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
            
            # Alternate row colors
            if i % 2 == 0:
                cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
        
        # Format result column with color
        result_cell = ws.cell(row=row, column=4)
        if result > 0:
            result_cell.font = Font(name='Calibri', size=9, bold=True, color="2E7D32")
            result_cell.value = f"${result:.2f}"
        elif result < 0:
            result_cell.font = Font(name='Calibri', size=9, bold=True, color="C62828")
            result_cell.value = f"${result:.2f}"
        else:
            result_cell.value = f"${result:.2f}"
        
        # Format status with color
        status_cell = ws.cell(row=row, column=5)
        if t.get('status') == 'open':
            status_cell.font = Font(name='Calibri', size=9, bold=True, color="FF6F00")
        
        row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 20
    ws.column_dimensions['H'].width = 15
    
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    
    # Generate filename with license key
    license_key_prefix = account_summary.get('license_key', 'unknown')[:8]
    filename = f'statement_{license_key_prefix}_{datetime.utcnow().date()}.xlsx'
    
    # Read the buffer content
    xlsx_data = bio.getvalue()
    
    # Return response with proper headers
    response = Response(xlsx_data, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.headers['Content-Length'] = str(len(xlsx_data))
    return response
